"""
views.py - Views REST e endpoints de exportação para Shopping Butantã / Quadro de Presenças.
"""

import io
import csv
import json
from typing import Any
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny

from .butanta import database
from .butanta.parser_engine import parse_raw_text


def sanitize_csv_cell(value: Any) -> str:
    """Protege contra ataques de Injeção de Fórmulas CSV."""
    text = str(value) if value is not None else ""
    if text and text[0] in ('=', '+', '-', '@', '\t', '\r'):
        return f"'{text}"
    return text


@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
def shoppings_view(request):
    """Listar ou adicionar shoppings."""
    if request.method == 'GET':
        try:
            shoppings = database.get_all_shoppings()
            return JsonResponse({"success": True, "shoppings": shoppings, "count": len(shoppings)})
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao listar shoppings: {str(e)}"}, status=500)

    elif request.method == 'POST':
        try:
            data = request.data if hasattr(request, 'data') else json.loads(request.body.decode('utf-8') or '{}')
            nome = str(data.get('nome', '')).strip()
            if not nome:
                return JsonResponse({"success": False, "error": "O nome do shopping é obrigatório."}, status=400)

            new_shopping = database.add_shopping(nome)
            return JsonResponse({
                "success": True,
                "message": f"Shopping '{new_shopping['nome']}' adicionado com sucesso!",
                "shopping": new_shopping
            }, status=201)
        except ValueError as ve:
            return JsonResponse({"success": False, "error": str(ve)}, status=409)
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao adicionar shopping: {str(e)}"}, status=500)


@api_view(['DELETE'])
@permission_classes([AllowAny])
def shopping_delete_view(request, shopping_id: int):
    """Retira/remove um shopping pelo ID e limpa seus registros vinculados."""
    try:
        result = database.delete_shopping(shopping_id)
        msg = f"Shopping '{result['nome']}' retirado com sucesso."
        if result['deleted_records'] > 0:
            msg += f" ({result['deleted_records']} registros vinculados foram excluídos)."
        return JsonResponse({"success": True, "message": msg, "result": result})
    except ValueError as ve:
        return JsonResponse({"success": False, "error": str(ve)}, status=404)
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Erro ao retirar shopping: {str(e)}"}, status=500)


@api_view(['POST'])
@permission_classes([AllowAny])
def parse_view(request):
    """
    Analisa o texto bruto de escala/presença.
    Retorna os registros identificados e destaca linhas não reconhecidas para revisão do usuário.
    """
    try:
        data = request.data if hasattr(request, 'data') else json.loads(request.body.decode('utf-8') or '{}')
        raw_text = str(data.get('text', '')).strip()
        auto_split = bool(data.get('auto_split', True))
        auto_date = bool(data.get('auto_date', True))
        shopping = str(data.get('shopping', '')).strip() or None

        if not raw_text:
            return JsonResponse({
                "success": False,
                "error": "Por favor, forneça o texto da mensagem para análise."
            }, status=400)

        result = parse_raw_text(raw_text, auto_split=auto_split, auto_date=auto_date, default_shopping=shopping)
        return JsonResponse(result)
    except Exception as e:
        return JsonResponse({
            "success": False,
            "error": f"Erro interno ao processar texto: {str(e)}"
        }, status=500)


@api_view(['GET', 'POST'])
@permission_classes([AllowAny])
def records_view(request):
    """Listar ou salvar registros de quadro."""
    if request.method == 'GET':
        try:
            month = request.GET.get('month', '').strip()
            search = request.GET.get('search', '').strip()
            turno = request.GET.get('turno', '').strip()
            shopping = request.GET.get('shopping', '').strip()

            records = database.get_all_records(
                month_filter=month if month else None,
                search=search if search else None,
                turno_filter=turno if turno else None,
                shopping_filter=shopping if shopping else None
            )
            return JsonResponse({"success": True, "records": records, "count": len(records)})
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao buscar registros: {str(e)}"}, status=500)

    elif request.method == 'POST':
        try:
            data = request.data if hasattr(request, 'data') else json.loads(request.body.decode('utf-8') or '{}')
            auto_merge = bool(data.get('auto_merge', True))
            
            raw_records = data.get('records') or ([data.get('record')] if data.get('record') else None)
            if not raw_records and isinstance(data, list):
                raw_records = data
            elif not raw_records and ('data' in data or 'date' in data):
                raw_records = [data]

            if not raw_records:
                return JsonResponse({"success": False, "error": "Nenhum registro fornecido para salvar."}, status=400)

            saved_results = []
            for rec in raw_records:
                if not isinstance(rec, dict):
                    continue
                saved_item = database.insert_or_merge_record(rec, auto_merge=auto_merge)
                saved_results.append(saved_item)

            return JsonResponse({
                "success": True,
                "message": f"{len(saved_results)} registro(s) processado(s) com sucesso.",
                "results": saved_results
            })
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao salvar registros: {str(e)}"}, status=500)


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([AllowAny])
def record_detail_view(request, record_id: int):
    """Obter, atualizar ou excluir um registro individual."""
    if request.method == 'GET':
        try:
            record = database.get_record_by_id(record_id)
            if not record:
                return JsonResponse({"success": False, "error": "Registro não encontrado."}, status=404)
            return JsonResponse({"success": True, "record": record})
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao carregar registro: {str(e)}"}, status=500)

    elif request.method in ('PUT', 'PATCH'):
        try:
            data = request.data if hasattr(request, 'data') else json.loads(request.body.decode('utf-8') or '{}')
            if not data:
                return JsonResponse({"success": False, "error": "Dados inválidos."}, status=400)

            success = database.update_record(record_id, data)
            if not success:
                return JsonResponse({"success": False, "error": "Registro não encontrado para atualização."}, status=404)

            return JsonResponse({"success": True, "message": "Registro atualizado com sucesso."})
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao atualizar registro: {str(e)}"}, status=500)

    elif request.method == 'DELETE':
        try:
            success = database.delete_record(record_id)
            if not success:
                return JsonResponse({"success": False, "error": "Registro não encontrado para exclusão."}, status=404)
            return JsonResponse({"success": True, "message": "Registro excluído com sucesso."})
        except Exception as e:
            return JsonResponse({"success": False, "error": f"Erro ao excluir registro: {str(e)}"}, status=500)


@api_view(['POST', 'DELETE'])
@permission_classes([AllowAny])
def records_clear_view(request):
    """Remove todos os registros do banco SQLite."""
    try:
        count = database.clear_all_records()
        return JsonResponse({"success": True, "message": f"Todos os {count} registros foram removidos com sucesso."})
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Erro ao limpar registros: {str(e)}"}, status=500)


@api_view(['GET'])
@permission_classes([AllowAny])
def summary_view(request):
    """Retorna os indicadores KPI e o resumo diário consolidado."""
    try:
        month = request.GET.get('month', '').strip()
        month_filter = month if month else None
        shopping = request.GET.get('shopping', '').strip()
        shopping_filter = shopping if shopping else None

        kpis = database.get_kpis(month_filter=month_filter, shopping_filter=shopping_filter)
        daily_summary = database.get_daily_summary(month_filter=month_filter, shopping_filter=shopping_filter)

        return JsonResponse({
            "success": True,
            "kpis": kpis,
            "daily_summary": daily_summary
        })
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Erro ao gerar resumo estatístico: {str(e)}"}, status=500)


@api_view(['GET'])
@permission_classes([AllowAny])
def export_csv_view(request):
    """Exporta arquivo CSV com UTF-8 BOM e separador ponto e vírgula (;)."""
    try:
        month = request.GET.get('month', '').strip()
        shopping = request.GET.get('shopping', '').strip()
        shopping_filter = shopping if shopping else None
        export_type = request.GET.get('type', 'daily').strip()

        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)

        headers = ['Data', 'Shopping', 'Turno', 'Presentes', 'Folgas', 'Faltas', 'Atestados', 'Apoio Noite', 'Banheirista Apoio', 'Outros', 'Observações']

        if export_type == 'detailed':
            records = database.get_all_records(
                month_filter=month if month else None,
                shopping_filter=shopping_filter
            )
            writer.writerow(headers)

            for r in records:
                outros_str = "; ".join(r.get('outros', []))
                writer.writerow([
                    sanitize_csv_cell(r.get('data_formatada', '')),
                    sanitize_csv_cell(r.get('shopping', '')),
                    sanitize_csv_cell(r.get('turno', '')),
                    r.get('presentes', 0),
                    r.get('folgas', 0),
                    r.get('faltas', 0),
                    r.get('atestados', 0),
                    r.get('apoio_noite', 0),
                    r.get('banheirista_apoio', 0),
                    sanitize_csv_cell(outros_str),
                    sanitize_csv_cell(r.get('observacoes', ''))
                ])
            filename = f"quadro_detalhado_{shopping or 'geral'}_{month or 'completo'}.csv"
        else:
            daily_rows = database.get_daily_summary(
                month_filter=month if month else None,
                shopping_filter=shopping_filter
            )
            writer.writerow(headers)

            tot_pres = tot_folg = tot_falt = tot_atest = tot_apoio = tot_banh = 0
            for r in daily_rows:
                tot_pres += r.get('presentes', 0)
                tot_folg += r.get('folgas', 0)
                tot_falt += r.get('faltas', 0)
                tot_atest += r.get('atestados', 0)
                tot_apoio += r.get('apoio_noite', 0)
                tot_banh += r.get('banheirista_apoio', 0)

                outros_str = "; ".join(r.get('outros', []))
                writer.writerow([
                    sanitize_csv_cell(r.get('data_formatada', '')),
                    sanitize_csv_cell(r.get('shopping', '')),
                    sanitize_csv_cell(r.get('turno', '')),
                    r.get('presentes', 0),
                    r.get('folgas', 0),
                    r.get('faltas', 0),
                    r.get('atestados', 0),
                    r.get('apoio_noite', 0),
                    r.get('banheirista_apoio', 0),
                    sanitize_csv_cell(outros_str),
                    sanitize_csv_cell(r.get('observacoes', ''))
                ])

            writer.writerow([])
            writer.writerow(['TOTAIS CONSOLIDADOS', '', '', tot_pres, tot_folg, tot_falt, tot_atest, tot_apoio, tot_banh, '', ''])
            filename = f"quadro_diario_{shopping or 'geral'}_{month or 'completo'}.csv"

        response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Erro ao exportar CSV: {str(e)}"}, status=500)


@api_view(['GET'])
@permission_classes([AllowAny])
def export_excel_view(request):
    """Gera uma planilha Excel estilizada e profissional (.xlsx) com duas abas formatadas."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        month = request.GET.get('month', '').strip()
        month_filter = month if month else None
        shopping = request.GET.get('shopping', '').strip()
        shopping_filter = shopping if shopping else None

        wb = openpyxl.Workbook()
        ws_daily = wb.active
        ws_daily.title = "Resumo Diário"
        ws_turnos = wb.create_sheet(title="Registros por Turno")

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        headers = ['Data', 'Shopping', 'Turno', 'Presentes', 'Folgas', 'Faltas', 'Atestados', 'Apoio Noite', 'Banheirista Apoio', 'Outros', 'Observações']

        # 1. Resumo Diário
        ws_daily.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws_daily.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        daily_rows = database.get_daily_summary(month_filter=month_filter, shopping_filter=shopping_filter)
        tot_pres = tot_folg = tot_falt = tot_atest = tot_apoio = tot_banh = 0

        for r in daily_rows:
            pres = r.get('presentes', 0)
            folg = r.get('folgas', 0)
            falt = r.get('faltas', 0)
            atest = r.get('atestados', 0)
            apoio = r.get('apoio_noite', 0)
            banh = r.get('banheirista_apoio', 0)

            tot_pres += pres
            tot_folg += folg
            tot_falt += falt
            tot_atest += atest
            tot_apoio += apoio
            tot_banh += banh

            outros_str = ", ".join(r.get('outros', []))
            row_data = [
                r.get('data_formatada', ''),
                r.get('shopping', 'Shopping Butantã'),
                'Dia Completo',
                pres, folg, falt, atest, apoio, banh,
                outros_str,
                r.get('observacoes', '')
            ]
            ws_daily.append(row_data)

        total_row_idx = ws_daily.max_row + 1
        ws_daily.append(['TOTAIS GERAIS', '', '', tot_pres, tot_folg, tot_falt, tot_atest, tot_apoio, tot_banh, '', ''])
        for col_num in range(1, len(headers) + 1):
            cell = ws_daily.cell(row=total_row_idx, column=col_num)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border
            if col_num in [1, 2, 3]:
                cell.alignment = align_left
            else:
                cell.alignment = align_center

        # 2. Registros por Turno
        ws_turnos.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws_turnos.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        all_records = database.get_all_records(month_filter=month_filter, shopping_filter=shopping_filter)
        for r in all_records:
            outros_str = ", ".join(r.get('outros', []))
            ws_turnos.append([
                r.get('data_formatada', ''),
                r.get('shopping', 'Shopping Butantã'),
                r.get('turno', ''),
                r.get('presentes', 0),
                r.get('folgas', 0),
                r.get('faltas', 0),
                r.get('atestados', 0),
                r.get('apoio_noite', 0),
                r.get('banheirista_apoio', 0),
                outros_str,
                r.get('observacoes', '')
            ])

        for ws in [ws_daily, ws_turnos]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row > 1:
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = align_center
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        mem = io.BytesIO()
        wb.save(mem)
        mem.seek(0)

        filename = f"quadro_presencas_{shopping or 'geral'}_{month or 'completo'}.xlsx"
        response = HttpResponse(
            mem.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        return JsonResponse({"success": False, "error": f"Erro ao exportar Excel: {str(e)}"}, status=500)
