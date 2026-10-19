from datetime import datetime, date, timedelta
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from usuarios.permissions import IsGestaoOrAdministrador
from ..models import Loja


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def loja_presencas_calendario_api(request, loja_id):
    """
    Por que existe: Retorna o consolidado de presenças registradas por dia para uma loja física
    específica em um determinado mês/ano.
    Evita carregar a lista de colaboradores de todos os dias de uma vez no frontend.
    """
    ano_mes = request.GET.get("ano_mes")  # Formato esperado: YYYY-MM (ex: 2026-05)
    if not ano_mes:
        return Response({"error": "Parâmetro 'ano_mes' é obrigatório (formato YYYY-MM)."}, status=status.HTTP_400_BAD_REQUEST)
    
    try:
        ano, mes = map(int, ano_mes.split("-"))
    except ValueError:
        return Response({"error": "Formato inválido de 'ano_mes'. Use YYYY-MM."}, status=status.HTTP_400_BAD_REQUEST)

    # Consulta consolidada agrupando por data
    from django.db.models import Count
    from colaboradores.models import PresencaRelogio

    dados = (
        PresencaRelogio.objects.filter(loja_id=loja_id, data__year=ano, data__month=mes)
        .values("data")
        .annotate(total=Count("id"))
        .order_by("data")
    )

    resultado = {d["data"].strftime("%Y-%m-%d"): d["total"] for d in dados}
    return Response(resultado, status=status.HTTP_200_OK)


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def loja_presencas_dia_api(request, loja_id):
    """
    Por que existe: Retorna o detalhamento (quem marcou presença e que horas)
    em uma determinada loja física em um dia específico.
    Isso é executado somente quando o usuário clica em um dia no calendário do frontend.
    """
    data_str = request.GET.get("data")  # Formato esperado: YYYY-MM-DD
    if not data_str:
        return Response({"error": "Parâmetro 'data' é obrigatório (formato YYYY-MM-DD)."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        data_obj = datetime.strptime(data_str, "%Y-%m-%d").date()
    except ValueError:
        return Response({"error": "Formato inválido de 'data'. Use YYYY-MM-DD."}, status=status.HTTP_400_BAD_REQUEST)

    from colaboradores.models import PresencaRelogio

    presencas = (
        PresencaRelogio.objects.filter(loja_id=loja_id, data=data_obj)
        .select_related("colaborador")
        .order_by("data_hora")
    )

    from django.utils import timezone

    lista = []
    for p in presencas:
        # Converte o datetime armazenado em UTC no banco para o fuso horário local configurado
        horario_local = timezone.localtime(p.data_hora)
        lista.append({
            "punch_id": p.punch_id,
            "cpf": p.cpf_original,
            "nome": p.colaborador.nome if p.colaborador else "Colaborador Não Identificado",
            "re": p.colaborador.re if p.colaborador else "N/A",
            "cargo": p.colaborador.cargo if p.colaborador else "N/A",
            "horario_entrada": horario_local.strftime("%H:%M:%S")
        })

    return Response(lista, status=status.HTTP_200_OK)


@api_view(["POST"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def lojas_presencas_sincronizar_geral_api(request):
    """
    Por que existe: Sincroniza em tempo real as batidas do dia anterior (ontem) de todas as lojas.
    """
    from colaboradores.services.geovictoria_punches_sync import sincronizar_punches_api
    
    ontem = date.today() - timedelta(days=1)

    try:
        res = sincronizar_punches_api(ontem, ontem, loja_id=None)
        return Response({
            "success": True,
            "message": "Sincronização geral de ontem concluída com sucesso!",
            "dados": res
        }, status=status.HTTP_200_OK)
    except Exception as e:
        return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(["POST"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def loja_presencas_sincronizar_individual_api(request, loja_id):
    """
    Por que existe: Sincroniza em tempo real as batidas dos últimos N dias (padrão: 30 dias) para uma loja específica.
    """
    from colaboradores.services.geovictoria_punches_sync import sincronizar_punches_api
    
    try:
        dias = int(request.data.get("dias") or request.GET.get("dias") or 30)
    except (ValueError, TypeError):
        dias = 30

    fim = date.today()
    inicio = fim - timedelta(days=dias)

    try:
        res = sincronizar_punches_api(inicio, fim, loja_id=loja_id)
        return Response({
            "success": True,
            "message": f"Sincronização individual dos últimos {dias} dias concluída com sucesso!",
            "dados": res
        }, status=status.HTTP_200_OK)
    except Exception as e:
        return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def loja_presencas_sincronizar_progresso_api(request):
    """
    Por que existe: Retorna o progresso em tempo real da sincronização de batidas do relógio GeoVictoria.
    """
    from django.core.cache import cache
    status_sync = cache.get("sync_punches_status")
    return Response(status_sync or {"page": 0, "total_pages": 0, "msg": "Aguardando..."}, status=status.HTTP_200_OK)


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsGestaoOrAdministrador])
def loja_presencas_exportar_excel_api(request, loja_id):
    """
    Exporta a auditoria de presenças da filial em planilha Excel oficial (.xlsx) com duas abas:
    - Aba 1: "Visão Total por Dia" (consolidado de cada dia do mês: data, dia da semana, total de presenças, percentual e situação).
    - Aba 2: "Presenças por Colaborador" (relação de cada pessoa, RE, cargo, CPF, total de presenças, dias com presença registrada e grade completa diária).
    """
    import calendar
    import re
    from io import BytesIO
    from datetime import date
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse
    from django.shortcuts import get_object_or_404
    from django.utils import timezone
    from colaboradores.models import PresencaRelogio

    loja = get_object_or_404(Loja, pk=loja_id)
    ano_mes = request.GET.get("ano_mes")
    if not ano_mes:
        hoje = date.today()
        ano, mes = hoje.year, hoje.month
    else:
        try:
            ano, mes = map(int, ano_mes.split("-"))
        except ValueError:
            return Response({"error": "Formato inválido de 'ano_mes'. Use YYYY-MM."}, status=status.HTTP_400_BAD_REQUEST)

    punches = (
        PresencaRelogio.objects.filter(loja_id=loja_id, data__year=ano, data__month=mes)
        .select_related("colaborador")
        .order_by("data", "data_hora")
    )

    _, num_days = calendar.monthrange(ano, mes)
    dias_semana_pt = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]

    diario_colabs = {d: set() for d in range(1, num_days + 1)}
    colabs_map = {}

    for p in punches:
        dia_num = p.data.day
        colab_key = (
            (p.colaborador.re if p.colaborador and p.colaborador.re else "") or
            p.cpf_original or
            (p.colaborador.nome if p.colaborador else "")
        )

        if colab_key not in colabs_map:
            colabs_map[colab_key] = {
                "re": p.colaborador.re if (p.colaborador and p.colaborador.re) else "N/A",
                "nome": p.colaborador.nome if p.colaborador else "Colaborador Não Identificado",
                "cargo": p.colaborador.cargo if (p.colaborador and p.colaborador.cargo) else "N/A",
                "cpf": p.cpf_original or "",
                "dias": {}
            }

        hora_local = timezone.localtime(p.data_hora).strftime("%H:%M")
        if dia_num not in colabs_map[colab_key]["dias"]:
            colabs_map[colab_key]["dias"][dia_num] = hora_local

        diario_colabs[dia_num].add(colab_key)

    wb = openpyxl.Workbook()
    ws_padrao = wb.active

    # Paleta de Estilos Corporativos
    font_titulo = Font(name="Calibri", size=14, bold=True, color="1E293B")
    font_subtitulo = Font(name="Calibri", size=10, bold=False, color="64748B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header_accent = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_presente = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_presente = Font(name="Calibri", size=9, bold=True, color="166534")
    font_total = Font(name="Calibri", size=10, bold=True, color="0F172A")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    border_total = Border(
        top=Side(style='thin', color='94A3B8'),
        bottom=Side(style='double', color='0F172A')
    )

    nome_loja_display = loja.nome_referencia or loja.nome_geovictoria or f"Loja {loja.id}"
    total_presencas_mes = sum(len(s) for s in diario_colabs.values())

    # ==========================================
    # ABA 1: VISÃO TOTAL DE CADA DIA
    # ==========================================
    ws1 = wb.create_sheet(title="Visão Total por Dia")
    if ws_padrao:
        wb.remove(ws_padrao)

    ws1["A1"] = f"Relatório de Presenças por Dia — {nome_loja_display}"
    ws1["A1"].font = font_titulo
    ws1["A2"] = f"Competência: {mes:02d}/{ano}  |  Centro de Custo: {loja.centro_de_custo or '-'}  |  Cliente: {loja.cliente or '-'}"
    ws1["A2"].font = font_subtitulo

    headers_aba1 = ["Data", "Dia da Semana", "Presenças Registradas", "% do Mês", "Situação"]
    for col_idx, h in enumerate(headers_aba1, start=1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for d in range(1, num_days + 1):
        dt = date(ano, mes, d)
        dia_sem = dias_semana_pt[dt.weekday()]
        qtd = len(diario_colabs[d])
        pct = (qtd / total_presencas_mes) if total_presencas_mes > 0 else 0
        situacao = "Com Presenças" if qtd > 0 else ("Final de Semana" if dt.weekday() >= 5 else "Sem Registro")

        row_num = 4 + d
        ws1.cell(row=row_num, column=1, value=dt.strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws1.cell(row=row_num, column=2, value=dia_sem).alignment = Alignment(horizontal="left")

        cell_qtd = ws1.cell(row=row_num, column=3, value=qtd)
        cell_qtd.alignment = Alignment(horizontal="center")
        cell_qtd.font = Font(name="Calibri", size=10, bold=qtd > 0)
        if qtd > 0:
            cell_qtd.fill = fill_presente

        cell_pct = ws1.cell(row=row_num, column=4, value=pct)
        cell_pct.alignment = Alignment(horizontal="center")
        cell_pct.number_format = "0.0%"

        ws1.cell(row=row_num, column=5, value=situacao).alignment = Alignment(horizontal="left")

        for c in range(1, 6):
            ws1.cell(row=row_num, column=c).border = thin_border
            if d % 2 == 0 and qtd == 0:
                ws1.cell(row=row_num, column=c).fill = fill_zebra

    tot_row = 4 + num_days + 1
    ws1.cell(row=tot_row, column=1, value="TOTAL GERAL").font = font_total
    ws1.cell(row=tot_row, column=2, value=f"{num_days} dias apurados").font = font_total
    cell_tot_qtd = ws1.cell(row=tot_row, column=3, value=total_presencas_mes)
    cell_tot_qtd.font = font_total
    cell_tot_qtd.alignment = Alignment(horizontal="center")
    cell_tot_pct = ws1.cell(row=tot_row, column=4, value=1.0)
    cell_tot_pct.font = font_total
    cell_tot_pct.alignment = Alignment(horizontal="center")
    cell_tot_pct.number_format = "0.0%"
    media_dia = total_presencas_mes / num_days if num_days > 0 else 0
    ws1.cell(row=tot_row, column=5, value=f"Média: {media_dia:.1f} presenças/dia").font = font_total

    for c in range(1, 6):
        cell = ws1.cell(row=tot_row, column=c)
        cell.fill = fill_total
        cell.border = border_total

    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # ==========================================
    # ABA 2: VISÃO POR PESSOA (COLABORADORES)
    # ==========================================
    ws2 = wb.create_sheet(title="Presenças por Colaborador")

    ws2["A1"] = f"Presenças Individuais — {nome_loja_display}"
    ws2["A1"].font = font_titulo
    ws2["A2"] = f"Competência: {mes:02d}/{ano}  |  Colaboradores com Registro: {len(colabs_map)}  |  Total de Presenças: {total_presencas_mes}"
    ws2["A2"].font = font_subtitulo

    headers_fixos = ["RE", "Colaborador", "Cargo", "CPF", "Total Presenças", "Dias com Presença Registrada"]
    headers_dias = [f"Dia {d:02d}" for d in range(1, num_days + 1)]
    headers_aba2 = headers_fixos + headers_dias

    for col_idx, h in enumerate(headers_aba2, start=1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header if col_idx <= len(headers_fixos) else fill_header_accent
        cell.alignment = Alignment(horizontal="center", vertical="center")

    colabs_ordenados = sorted(colabs_map.values(), key=lambda x: str(x["nome"]).strip().upper())

    for idx, c_info in enumerate(colabs_ordenados, start=1):
        row_num = 4 + idx
        dias_dict = c_info["dias"]
        qtd_presencas = len(dias_dict)
        dias_str = ", ".join(f"{d:02d}" for d in sorted(dias_dict.keys()))

        ws2.cell(row=row_num, column=1, value=c_info["re"]).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_num, column=2, value=c_info["nome"]).alignment = Alignment(horizontal="left")
        ws2.cell(row=row_num, column=3, value=c_info["cargo"]).alignment = Alignment(horizontal="left")
        ws2.cell(row=row_num, column=4, value=c_info["cpf"]).alignment = Alignment(horizontal="center")

        cell_tot_p = ws2.cell(row=row_num, column=5, value=qtd_presencas)
        cell_tot_p.alignment = Alignment(horizontal="center")
        cell_tot_p.font = Font(name="Calibri", size=10, bold=True)

        ws2.cell(row=row_num, column=6, value=dias_str).alignment = Alignment(horizontal="left")

        for d in range(1, num_days + 1):
            col_dia = len(headers_fixos) + d
            cell_dia = ws2.cell(row=row_num, column=col_dia)
            if d in dias_dict:
                cell_dia.value = dias_dict[d]
                cell_dia.fill = fill_presente
                cell_dia.font = font_presente
                cell_dia.alignment = Alignment(horizontal="center")
            else:
                cell_dia.value = "-"
                cell_dia.alignment = Alignment(horizontal="center")
                cell_dia.font = Font(name="Calibri", size=9, color="94A3B8")

        for c in range(1, len(headers_aba2) + 1):
            ws2.cell(row=row_num, column=c).border = thin_border
            if idx % 2 == 0 and c <= len(headers_fixos):
                ws2.cell(row=row_num, column=c).fill = fill_zebra

    row_tot2 = 4 + len(colabs_ordenados) + 1
    ws2.cell(row=row_tot2, column=1, value="TOTAL").font = font_total
    ws2.cell(row=row_tot2, column=2, value=f"{len(colabs_ordenados)} colaboradores").font = font_total
    ws2.cell(row=row_tot2, column=3, value="-").font = font_total
    ws2.cell(row=row_tot2, column=4, value="-").font = font_total
    cell_soma2 = ws2.cell(row=row_tot2, column=5, value=total_presencas_mes)
    cell_soma2.font = font_total
    cell_soma2.alignment = Alignment(horizontal="center")
    ws2.cell(row=row_tot2, column=6, value="-").font = font_total

    for d in range(1, num_days + 1):
        col_dia = len(headers_fixos) + d
        qtd_dia = len(diario_colabs[d])
        cell_tot_dia = ws2.cell(row=row_tot2, column=col_dia, value=qtd_dia)
        cell_tot_dia.font = font_total
        cell_tot_dia.alignment = Alignment(horizontal="center")

    for c in range(1, len(headers_aba2) + 1):
        cell = ws2.cell(row=row_tot2, column=c)
        cell.fill = fill_total
        cell.border = border_total

    for col_idx in range(1, len(headers_fixos) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws2.cell(row=r, column=col_idx).value or "")) for r in range(4, row_tot2 + 1))
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    for d in range(1, num_days + 1):
        col_letter = get_column_letter(len(headers_fixos) + d)
        ws2.column_dimensions[col_letter].width = 9

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    clean_name = re.sub(r'[^a-zA-Z0-9_-]', '_', nome_loja_display.strip())
    filename = f"presencas_{clean_name}_{ano}_{mes:02d}.xlsx"

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
