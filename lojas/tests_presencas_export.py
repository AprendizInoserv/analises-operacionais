from datetime import date, datetime
from io import BytesIO
import openpyxl
from django.test import TestCase
from rest_framework.test import APIClient
from django.contrib.auth import get_user_model
from lojas.models import Loja
from colaboradores.models import Colaborador, PresencaRelogio

User = get_user_model()


class PresencasExportExcelTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="admin_test_export",
            email="admin_export@teste.com",
            password="password123"
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        self.loja = Loja.objects.create(
            nome_geovictoria="LOJA TESTE PRESENCAS 101",
            nome_referencia="LOJA TESTE 101",
            cliente="ATACADAO",
            centro_de_custo="CC101"
        )

        self.colab1 = Colaborador.objects.create(
            nome="ANA SILVA",
            re="10001",
            cargo="OPERADORA DE LIMPEZA",
            cpf="11122233344",
            data_admissao=date(2025, 1, 1)
        )
        self.colab2 = Colaborador.objects.create(
            nome="BRUNO SOUZA",
            re="10002",
            cargo="LIDER DE EQUIPE",
            cpf="55566677788",
            data_admissao=date(2025, 1, 1)
        )

        # Batidas de teste em outubro de 2026
        PresencaRelogio.objects.create(
            punch_id="P1",
            colaborador=self.colab1,
            cpf_original=self.colab1.cpf,
            loja=self.loja,
            grupo_geovictoria="LOJA TESTE PRESENCAS 101",
            data=date(2026, 10, 1),
            data_hora=datetime(2026, 10, 1, 8, 15, 0)
        )
        PresencaRelogio.objects.create(
            punch_id="P2",
            colaborador=self.colab1,
            cpf_original=self.colab1.cpf,
            loja=self.loja,
            grupo_geovictoria="LOJA TESTE PRESENCAS 101",
            data=date(2026, 10, 2),
            data_hora=datetime(2026, 10, 2, 8, 20, 0)
        )
        PresencaRelogio.objects.create(
            punch_id="P3",
            colaborador=self.colab2,
            cpf_original=self.colab2.cpf,
            loja=self.loja,
            grupo_geovictoria="LOJA TESTE PRESENCAS 101",
            data=date(2026, 10, 2),
            data_hora=datetime(2026, 10, 2, 8, 0, 0)
        )

    def test_exportar_excel_duas_abas(self):
        url = f"/lojas/api/presencas/exportar/{self.loja.id}/?ano_mes=2026-10"
        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("presencas_", response["Content-Disposition"])

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["Visão Total por Dia", "Presenças por Colaborador"])

        # Aba 1: Visão Total por Dia
        ws1 = wb["Visão Total por Dia"]
        self.assertIn("Relatório de Presenças por Dia", ws1["A1"].value)
        # Dia 1 (outubro): 1 presença (Ana)
        self.assertEqual(ws1.cell(5, 1).value, "01/10/2026")
        self.assertEqual(ws1.cell(5, 3).value, 1)
        # Dia 2 (outubro): 2 presenças (Ana e Bruno)
        self.assertEqual(ws1.cell(6, 1).value, "02/10/2026")
        self.assertEqual(ws1.cell(6, 3).value, 2)

        # Aba 2: Presenças por Colaborador
        ws2 = wb["Presenças por Colaborador"]
        self.assertIn("Presenças Individuais", ws2["A1"].value)
        # Ana Silva
        self.assertEqual(ws2.cell(5, 2).value, "ANA SILVA")
        self.assertEqual(ws2.cell(5, 5).value, 2)  # 2 dias presentes
        self.assertEqual(ws2.cell(5, 6).value, "01, 02")
        # Bruno Souza
        self.assertEqual(ws2.cell(6, 2).value, "BRUNO SOUZA")
        self.assertEqual(ws2.cell(6, 5).value, 1)  # 1 dia presente
        self.assertEqual(ws2.cell(6, 6).value, "02")
