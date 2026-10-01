from datetime import date
from django.contrib.auth.models import User
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework import status

from lojas.models import Loja, Coordenador
from colaboradores.models import Colaborador


class HeadcountViewsTests(TestCase):
    """
    Por que existe: Esta classe testa a lógica de cálculo de headcount consolidado por loja
    e a listagem nominal de colaboradores correspondente, incluindo a regra especial
    que permite contar o status FÉRIAS apenas para clientes do grupo ATACADÃO.
    Trabalha com lojas ativas, sem dependência de data e com paginação ativada.
    """

    def setUp(self):
        # Cria um usuário de teste e configura o cliente da API
        self.user = User.objects.create_superuser(username="testanalista", password="password123")
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        # Cria coordenador de teste
        self.coordenador = Coordenador.objects.create(nome="Coordenador Roberto")

        # Cria lojas ativas (uma Atacadão com coordenador, uma normal sem coordenador) com o quadro planejado no cadastro
        self.loja_atacadao = Loja.objects.create(
            nome_referencia="ATACADÃO SÃO PAULO",
            cliente="ATACADÃO",
            centro_de_custo="123456789012",
            quadro="3",
            uf="SP",
            status="ATIVA",
            coordenador=self.coordenador,
        )
        self.loja_carrefour = Loja.objects.create(
            nome_referencia="CARREFOUR CAMPINAS",
            cliente="CARREFOUR",
            centro_de_custo="123456789013",
            quadro="2",
            uf="SP",
            status="ATIVA",
        )
        # Cria uma loja inativa para validar que ela é excluída do relatório
        self.loja_inativa = Loja.objects.create(
            nome_referencia="LOJA INATIVA DE TESTE",
            cliente="TESTE",
            centro_de_custo="123456789014",
            quadro="5",
            uf="SP",
            status="INATIVA",
        )

        # Cria colaboradores no TOTVS SRA para o Atacadão (Quadro planejado 3: 3 ativos + 1 férias = 4 real, desvio +1)
        Colaborador.objects.create(
            re="1001",
            nome="Jose Ativo Atacadao",
            loja=self.loja_atacadao,
            data_admissao=date(2026, 1, 1),
            status="",
            cargo="OPERADOR",
        )
        Colaborador.objects.create(
            re="1002",
            nome="Maria Ativo 2 Atacadao",
            loja=self.loja_atacadao,
            data_admissao=date(2026, 1, 1),
            status="",
            cargo="OPERADOR",
        )
        Colaborador.objects.create(
            re="1003",
            nome="Joao Ferias Atacadao",
            loja=self.loja_atacadao,
            data_admissao=date(2026, 1, 1),
            status="F",
            cargo="OPERADOR",
        )
        Colaborador.objects.create(
            re="1004",
            nome="Carlos Ativo 3 Atacadao",
            loja=self.loja_atacadao,
            data_admissao=date(2026, 1, 1),
            status="",
            cargo="OPERADOR",
        )
        Colaborador.objects.create(
            re="1005",
            nome="Pedro Demitido Atacadao",
            loja=self.loja_atacadao,
            data_admissao=date(2026, 1, 1),
            status="D",
            cargo="OPERADOR",
        )

        # Cria colaboradores no TOTVS SRA para o Carrefour (Quadro planejado 2: 1 ativo = 1 real. Férias ignoradas para Carrefour, desvio -1)
        Colaborador.objects.create(
            re="2001",
            nome="Ana Ativo Carrefour",
            loja=self.loja_carrefour,
            data_admissao=date(2026, 1, 1),
            status="",
            cargo="AUXILIAR",
        )
        Colaborador.objects.create(
            re="2002",
            nome="Lucas Demitido Carrefour",
            loja=self.loja_carrefour,
            data_admissao=date(2026, 1, 1),
            status="D",
            cargo="AUXILIAR",
        )
        Colaborador.objects.create(
            re="2003",
            nome="Rita Ferias Carrefour",
            loja=self.loja_carrefour,
            data_admissao=date(2026, 1, 1),
            status="F",
            cargo="AUXILIAR",
        )

    def test_headcount_analise_consolidado(self):
        # Por que existe: Este teste valida que o endpoint retorna corretamente os KPIs globais,
        # calculando o headcount_real dinamicamente através do status TOTVS SRA,
        # incluindo o total_excedentes que soma apenas os excedentes individuais de cada loja (desvio > 0)
        # em vez de somar desvios negativos, e a paginação das lojas ativas.

        # Faz a chamada para a API consolidada (sem filtros de data)
        response = self.client.get("/lojas/headcount/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # Como a API agora é paginada, os dados estarão estruturados sob a chave results
        results_data = response.data["results"]
        resultados = results_data["resultados"]
        
        # Devem ser listadas apenas as 2 lojas ATIVAS
        self.assertEqual(len(resultados), 2)
        self.assertNotIn("LOJA INATIVA DE TESTE", [r["nome_referencia"] for r in resultados])

        # Valida dados do Atacadão (Planejado: 3, Real: 4, Desvio: 1)
        atacadao_data = next(
            r for r in resultados if r["nome_referencia"] == "ATACADÃO SÃO PAULO"
        )
        self.assertEqual(atacadao_data["quadro_planejado"], 3)
        self.assertEqual(atacadao_data["headcount_real"], 4)
        self.assertEqual(atacadao_data["desvio"], 1)
        self.assertTrue(atacadao_data["is_atacadao"])

        # Valida dados do Carrefour (Planejado: 2, Real: 1, Desvio: -1)
        carrefour_data = next(
            r for r in resultados if r["nome_referencia"] == "CARREFOUR CAMPINAS"
        )
        self.assertEqual(carrefour_data["quadro_planejado"], 2)
        self.assertEqual(carrefour_data["headcount_real"], 1)
        self.assertEqual(carrefour_data["desvio"], -1)
        self.assertFalse(carrefour_data["is_atacadao"])

        # Valida KPIs globais
        kpis = results_data["kpis"]
        self.assertEqual(kpis["total_planejado"], 5)  # 3 (Atacadão) + 2 (Carrefour)
        self.assertEqual(kpis["total_real"], 5)       # 4 (Atacadão) + 1 (Carrefour)
        self.assertEqual(kpis["total_excedentes"], 1)  # Apenas o excedente do Atacadão (+1)
        self.assertEqual(kpis["total_lojas"], 2)

    def test_headcount_loja_colaboradores_detail(self):
        # Valida listagem nominal do Atacadão (deve trazer os 4 válidos: Jose, Maria, Joao férias e Carlos)
        url_atacadao = f"/lojas/headcount/{self.loja_atacadao.id}/colaboradores/"
        response_atacadao = self.client.get(url_atacadao)
        self.assertEqual(response_atacadao.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response_atacadao.data), 4)

        nomes_atacadao = [c["nome"] for c in response_atacadao.data]
        self.assertIn("Jose Ativo Atacadao", nomes_atacadao)
        self.assertIn("Maria Ativo 2 Atacadao", nomes_atacadao)
        self.assertIn("Joao Ferias Atacadao", nomes_atacadao)
        self.assertIn("Carlos Ativo 3 Atacadao", nomes_atacadao)
        self.assertNotIn("Pedro Demitido Atacadao", nomes_atacadao)

        # Valida listagem nominal do Carrefour (deve trazer 1 válido: Ana. Lucas é demitido e Rita é férias)
        url_carrefour = f"/lojas/headcount/{self.loja_carrefour.id}/colaboradores/"
        response_carrefour = self.client.get(url_carrefour)
        self.assertEqual(response_carrefour.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response_carrefour.data), 1)

        nomes_carrefour = [c["nome"] for c in response_carrefour.data]
        self.assertIn("Ana Ativo Carrefour", nomes_carrefour)
        self.assertNotIn("Lucas Demitido Carrefour", nomes_carrefour)
        self.assertNotIn("Rita Ferias Carrefour", nomes_carrefour)

    def test_headcount_coordenador_e_busca(self):
        # Valida que o coordenador vem preenchido quando vinculado e "-" quando não vinculado
        response = self.client.get("/lojas/headcount/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        resultados = response.data["results"]["resultados"]

        atacadao_data = next(r for r in resultados if r["nome_referencia"] == "ATACADÃO SÃO PAULO")
        carrefour_data = next(r for r in resultados if r["nome_referencia"] == "CARREFOUR CAMPINAS")

        self.assertEqual(atacadao_data["coordenador"], "Coordenador Roberto")
        self.assertEqual(carrefour_data["coordenador"], "-")

        # Valida busca pelo nome do coordenador
        response_busca = self.client.get("/lojas/headcount/?busca=Roberto")
        self.assertEqual(response_busca.status_code, status.HTTP_200_OK)
        resultados_busca = response_busca.data["results"]["resultados"]
        self.assertEqual(len(resultados_busca), 1)
        self.assertEqual(resultados_busca[0]["nome_referencia"], "ATACADÃO SÃO PAULO")

    def test_headcount_exportar_excel(self):
        # Valida endpoint de exportação para Excel (.xlsx)
        response = self.client.get("/lojas/headcount/exportar/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

        # Lê o Excel gerado para validar colunas e conteúdo
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Headcount", wb.sheetnames)
        sheet = wb["Headcount"]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Coordenador", headers)
        self.assertIn("Loja", headers)
        self.assertIn("Cliente", headers)
        self.assertIn("Centro de Custo", headers)
        self.assertIn("Ativos TOTVS", headers)
        self.assertIn("Quadro Planejado", headers)
        self.assertIn("Desvio", headers)
        self.assertIn("Presenças Ontem", headers)

    def test_lojas_exportar_excel(self):
        # Configura nomes extras para testar exportação completa de nomes
        self.loja_atacadao.nome_totvs = "ATACADAO TOTVS 10"
        self.loja_atacadao.nome_geovictoria = "ATACADAO GEO 10"
        self.loja_atacadao.nome_financeiro = "ATACADAO FIN 10"
        self.loja_atacadao.nome_findme = "ATACADAO FINDME 10"
        self.loja_atacadao.nome_metricas = "ATACADAO METRICAS 10"
        self.loja_atacadao.codigo_loja = 101
        self.loja_atacadao.save()

        response = self.client.get("/lojas/exportar/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Lojas", wb.sheetnames)
        sheet = wb["Lojas"]
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Cód. Loja", headers)
        self.assertIn("Nome de Referência", headers)
        self.assertIn("Nome TOTVS", headers)
        self.assertIn("Nome GeoVictoria", headers)
        self.assertIn("Nome Financeiro", headers)
        self.assertIn("Nome FindMe", headers)
        self.assertIn("Nome Métricas", headers)
        self.assertIn("Cliente/Regional", headers)
        self.assertIn("Centro de Custo", headers)
        self.assertIn("Quadro Estimado", headers)
        self.assertIn("Coordenador", headers)
        self.assertIn("Supervisor", headers)
        self.assertIn("Status", headers)
