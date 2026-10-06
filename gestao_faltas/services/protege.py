import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

MESES_NOMES = [
    '', 'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
    'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro'
]

def exportar_excel_protege(fechamento):
    """
    Gera o arquivo Excel no layout oficial do fechamento da Protege (faltas_protege.xlsx).
    """
    wb = Workbook()
    ws = wb.active
    
    mes_num = fechamento.mes_referencia
    nome_mes = MESES_NOMES[mes_num] if 1 <= mes_num <= 12 else str(mes_num)
    ws.title = nome_mes

    # Estilos
    font_header = Font(name='Calibri', size=11, bold=True, color='000000')
    font_data = Font(name='Calibri', size=11, color='000000')
    font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
    fill_header = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='CBD5E1')
    border_cell = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    headers = ['LOJA', 'FALTAS E ATESTADOS', 'CARGO', 'DATA', 'CONTATO', 'EMAIL´S DE CONTAOS']
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center' if col_idx in [2, 4] else 'left', vertical='center')
        cell.border = border_cell

    itens = fechamento.itens.select_related('loja').order_by('ordem', 'loja__ordem', 'loja__nome')
    row_num = 2
    for item in itens:
        loja = item.loja
        faltas_val = item.faltas_computadas
        cargo_val = item.cargo or ('SEM FALTAS' if faltas_val == 0 else 'AUX. DE LIMPEZA')
        data_val = item.status_data or ''
        contato_val = loja.supervisor or ''
        emails_val = loja.email_contatos or ''

        row_data = [
            loja.nome,
            faltas_val,
            cargo_val,
            data_val,
            contato_val,
            emails_val
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = font_bold if col_idx == 1 else font_data
            cell.border = border_cell
            if col_idx == 2:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        row_num += 1

    # Linha 21 (ou 3 linhas após o último registro) para o corpo e assunto do e-mail
    email_row = max(row_num + 3, 21)

    corpo_email = (
        "Boa tarde,\n\n"
        "Segue o fechamento das faltas apuradas para validação.\n\n"
        "Caso identifiquem alguma divergência, solicitamos a gentileza de nos informar "
        "o mais breve possível, para que as devidas retificações possam ser realizadas antes do faturamento.\n\n"
        "Não havendo manifestação, o processo será considerado validado e a nota fiscal será "
        "encaminhada para faturamento em até 24 horas.\n\n"
        "Atenciosamente,"
    )
    assunto_email = f"Fechamento de faltas - Protege/{nome_mes}"

    ws.cell(row=email_row, column=2, value=corpo_email)
    ws.cell(row=email_row, column=2).font = Font(name='Calibri', size=10, italic=True)
    ws.cell(row=email_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')

    ws.cell(row=email_row, column=3, value=assunto_email)
    ws.cell(row=email_row, column=3).font = Font(name='Calibri', size=11, bold=True)
    ws.cell(row=email_row, column=3).alignment = Alignment(vertical='top')

    # Ajuste de larguras das colunas
    col_widths = {
        'A': 30,
        'B': 24,
        'C': 26,
        'D': 16,
        'E': 32,
        'F': 85
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
