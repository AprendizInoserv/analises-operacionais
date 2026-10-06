import io
from datetime import timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from gestao_faltas.models import LojaTarifa as Loja

MESES_NOMES = [
    '', 'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
    'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro'
]

DIAS_SEMANA = ['Seg', 'Ter', 'Qua', 'Qui', 'Sex', 'Sáb', 'Dom']

def obter_dias_periodo(data_inicio, data_fim):
    """Retorna lista de objetos date do início ao fim do período."""
    dias = []
    curr = data_inicio
    while curr <= data_fim:
        dias.append(curr)
        curr += timedelta(days=1)
    return dias

def gerar_matriz_detalhe(fechamento, loja_id=None):
    """
    Gera a matriz completa de datas repetidas e colaboradores conforme fluxo do Exemplocarrefour.pdf.
    Cada dia do ciclo (do dia 20 ao dia 19) se repete N vezes por loja,
    onde N é a quantidade de colaboradores ativos daquela loja.
    """
    data_inicio = fechamento.data_inicio
    data_fim = fechamento.data_fim
    dias = obter_dias_periodo(data_inicio, data_fim)
    total_dias = len(dias)

    # Filtrar lojas
    qs_lojas = Loja.objects.filter(cliente=fechamento.cliente, ativo=True)
    if loja_id:
        qs_lojas = qs_lojas.filter(id=loja_id)
    
    lojas = list(qs_lojas.select_related('regiao').prefetch_related('colaboradores').order_by('regiao__ordem', 'ordem', 'nome'))

    linhas = []
    resumo_lojas = []
    total_colaboradores_geral = 0

    for loja in lojas:
        colabs = list(loja.colaboradores.filter(status='ATIVO').order_by('nome'))
        qtd_colabs = len(colabs)
        total_colaboradores_geral += qtd_colabs
        total_linhas_loja = qtd_colabs * total_dias

        resumo_lojas.append({
            'loja_id': loja.id,
            'loja_nome': loja.nome,
            'sigla': loja.sigla or '',
            'centro': loja.centro or loja.codigo_filial or '',
            'regional': loja.regiao.nome,
            'formato': loja.formato or 'HIPER',
            'desconto_unitario': float(loja.desconto_por_falta),
            'total_colaboradores': qtd_colabs,
            'repeticoes_por_data': qtd_colabs,
            'total_linhas': total_linhas_loja,
            'colaboradores_nomes': [c.nome for c in colabs]
        })

        if qtd_colabs == 0:
            continue

        # Para cada dia do período:
        # Repete a data para cada colaborador da loja
        for d in dias:
            data_formatada = d.strftime('%d/%m/%Y')
            data_iso = d.strftime('%Y-%m-%d')
            dia_sem = DIAS_SEMANA[d.weekday()]
            mes_nome = MESES_NOMES[d.month] if 1 <= d.month <= 12 else ''

            for colab in colabs:
                linhas.append({
                    'loja_id': loja.id,
                    'loja_nome': loja.nome,
                    'sigla': loja.sigla or '',
                    'centro': loja.centro or loja.codigo_filial or '',
                    'regional': loja.regiao.nome,
                    'formato': loja.formato or 'HIPER',
                    'colaborador_id': colab.id,
                    'colaborador_nome': colab.nome,
                    'turno': colab.turno or '',
                    'horario': colab.horario or '',
                    'dia': data_formatada,            # Coluna M do Carrefour
                    'dia_iso': data_iso,
                    'dia_semana': dia_sem,
                    'presente': 'Presente',           # Coluna N
                    'quem_cobriu': '',                # Coluna O
                    'observacoes': 'Presente',        # Coluna P
                    'status': 'Presente',             # Coluna Q
                    'desconto': 0.00,                 # Coluna R
                    'mes': mes_nome,                  # Coluna S
                })

    # Strings pré-formatadas para cópia direta
    coluna_m_datas = "\n".join(r['dia'] for r in linhas)
    coluna_j_nomes = "\n".join(r['colaborador_nome'] for r in linhas)

    # TSV completo para colar direto no Excel (a partir da coluna E)
    tsv_linhas = []
    cabecalho_tsv = "Nome da loja\tSigla\tCentro\tRegional\tFormato\tNome do funcionário\tTurno\tHorário\tDia\tPresente\tQuem cobriu falta\tObservações\tStatus\tDesconto\tMês"
    tsv_linhas.append(cabecalho_tsv)
    for r in linhas:
        linha_str = (
            f"{r['loja_nome']}\t{r['sigla']}\t{r['centro']}\t{r['regional']}\t{r['formato']}\t"
            f"{r['colaborador_nome']}\t{r['turno']}\t{r['horario']}\t{r['dia']}\t{r['presente']}\t"
            f"{r['quem_cobriu']}\t{r['observacoes']}\t{r['status']}\t{r['desconto']:.2f}\t{r['mes']}"
        )
        tsv_linhas.append(linha_str)
    tsv_completo = "\n".join(tsv_linhas)

    loja_selecionada_info = None
    if loja_id and resumo_lojas:
        loja_selecionada_info = resumo_lojas[0]

    return {
        'fechamento_id': fechamento.id,
        'fechamento_titulo': fechamento.titulo,
        'cliente': fechamento.cliente.nome,
        'periodo': f"{data_inicio.strftime('%d/%m/%Y')} até {data_fim.strftime('%d/%m/%Y')}",
        'data_inicio': data_inicio.strftime('%Y-%m-%d'),
        'data_fim': data_fim.strftime('%Y-%m-%d'),
        'total_dias': total_dias,
        'total_lojas': len(resumo_lojas),
        'total_colaboradores': total_colaboradores_geral,
        'total_linhas': len(linhas),
        'loja_selecionada': loja_selecionada_info,
        'resumo_lojas': resumo_lojas,
        'linhas': linhas,
        'coluna_m_datas': coluna_m_datas,
        'coluna_j_nomes': coluna_j_nomes,
        'tsv_completo': tsv_completo
    }


def exportar_excel_detalhe(fechamento, loja_id=None):
    """
    Gera um arquivo Excel .xlsx com a aba 'Detalhe' exatamente estruturada
    como no modelo oficial do Grupo Carrefour Brasil (colunas E a S).
    """
    dados = gerar_matriz_detalhe(fechamento, loja_id=loja_id)
    linhas = dados['linhas']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalhe"
    ws.views.sheetView[0].showGridLines = True

    # Cores
    header_fill = PatternFill(start_color="1B5E8A", end_color="1B5E8A", fill_type="solid")
    title_fill = PatternFill(start_color="15486B", end_color="15486B", fill_type="solid")
    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")

    font_title = Font(name="Aptos Narrow", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="Aptos Narrow", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Aptos Narrow", size=10, bold=False, color="1A202C")
    font_bold = Font(name="Aptos Narrow", size=10, bold=True, color="1A202C")

    thin_border_side = Side(border_style="thin", color="CBD5E0")
    thin_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side
    )

    # Linha 1: Título Carrefour
    ws.merge_cells("B1:S1")
    title_cell = ws["B1"]
    title_cell.value = f"Painel de controle de faltas - Grupo Carrefour Brasil | {dados['fechamento_titulo']}"
    title_cell.font = font_title
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Linha 2: Cabeçalho das Colunas (Colunas E a S)
    headers = {
        'C': ':: Painel ::',
        'E': 'Nome da loja',
        'F': 'Sigla',
        'G': 'Centro',
        'H': 'Regional',
        'I': 'Formato',
        'J': 'Nome do funcionário',
        'K': 'Turno',
        'L': 'Horário',
        'M': 'Dia',
        'N': 'Presente',
        'O': 'Quem cobriu falta',
        'P': 'Observações',
        'Q': 'Status',
        'R': 'Desconto',
        'S': 'Mês'
    }

    ws.row_dimensions[2].height = 24
    for col_letter, title in headers.items():
        cell = ws[f"{col_letter}2"]
        cell.value = title
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Inserção das Linhas de Dados
    current_row = 3
    for idx, r in enumerate(linhas):
        ws.row_dimensions[current_row].height = 19
        use_zebra = (idx % 2 == 1)
        row_fill = zebra_fill if use_zebra else PatternFill(fill_type=None)

        cells_map = {
            'E': (r['loja_nome'], 'left', font_bold),
            'F': (r['sigla'], 'center', font_data),
            'G': (r['centro'], 'center', font_data),
            'H': (r['regional'], 'left', font_data),
            'I': (r['formato'], 'center', font_data),
            'J': (r['colaborador_nome'], 'left', font_data),
            'K': (r['turno'], 'center', font_data),
            'L': (r['horario'], 'center', font_data),
            'M': (r['dia'], 'center', font_bold),     # Coluna M: Dia
            'N': (r['presente'], 'center', font_data),
            'O': (r['quem_cobriu'], 'left', font_data),
            'P': (r['observacoes'], 'left', font_data),
            'Q': (r['status'], 'center', font_data),
            'R': (r['desconto'], 'right', font_data),
            'S': (r['mes'], 'center', font_data),
        }

        for col_letter, (val, align, font_to_use) in cells_map.items():
            c = ws[f"{col_letter}{current_row}"]
            c.value = val
            c.font = font_to_use
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = thin_border
            if use_zebra:
                c.fill = row_fill
            if col_letter == 'R':
                c.number_format = 'R$ #,##0.00'

        current_row += 1

    # Larguras das colunas
    col_widths = {
        'A': 3, 'B': 5, 'C': 14, 'D': 3,
        'E': 26, 'F': 8, 'G': 10, 'H': 24, 'I': 12,
        'J': 36, 'K': 10, 'L': 14, 'M': 14,
        'N': 12, 'O': 22, 'P': 22, 'Q': 14,
        'R': 14, 'S': 12
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
